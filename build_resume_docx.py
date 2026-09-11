"""
build_resume_docx.py —— 简历 Markdown → Word(.docx)
====================================================
把 `桌面/简历-改写版2.md` 排成可直接投递的 Word 文件。

设计约束（**为 ATS 友好**，不是为好看）：
  - **不用表格、不用文本框、不用分栏**——大量 ATS 解析不了表格里的内容，
    一进简历库就丢字段。全部用普通段落 + Word 内置列表样式
  - 用**内置样式**（List Bullet / List Bullet 2）而不是手打「•」——
    解析器认样式，认不出手打符号的层级
  - 字体同时设置 ascii 与 eastAsia（中文必须设 w:eastAsia，否则 Word 回落默认字体）

用法:
    python build_resume_docx.py                    # 默认读桌面那份 md
    python build_resume_docx.py 输入.md 输出.docx
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT_CN = "微软雅黑"
FONT_EN = "微软雅黑"
GREY = RGBColor(0x66, 0x66, 0x66)
BLUE = RGBColor(0x1F, 0x4E, 0x79)

DEFAULT_SRC = Path.home() / "Desktop" / "简历-改写版2.md"
DEFAULT_OUT = Path.home() / "Desktop" / "朱科宇-AI应用开发-简历.docx"


def _style_run(run, size=10, bold=False, color=None, italic=False):
    """中英文字体都要显式设——只设 font.name 时中文会回落默认字体。"""
    run.font.name = FONT_EN
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)


def _add_runs(par, text, size=10, base_bold=False, color=None):
    """把 `**粗**` / `*斜*` 拆成多个 run，保留行内强调。"""
    for seg in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**"):
            _style_run(par.add_run(seg[2:-2]), size=size, bold=True,
                       color=color, italic=(color is GREY))
        elif seg.startswith("*") and seg.endswith("*") and len(seg) > 2:
            _style_run(par.add_run(seg[1:-1]), size=size, bold=base_bold,
                       color=color, italic=True)
        else:
            _style_run(par.add_run(seg), size=size, bold=base_bold, color=color)


def _bottom_border(par, color="BFBFBF", sz="6"):
    pPr = par._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def _tight(par, before=0, after=2, line=None):
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = line


def build(src: Path, out: Path) -> None:
    doc = Document()

    # 页面：A4 + 适中页边距（窄边距塞得下更多，但别窄到打印机裁切）
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for attr, val in (("top_margin", 1.4), ("bottom_margin", 1.4),
                      ("left_margin", 1.6), ("right_margin", 1.6)):
        setattr(sec, attr, Cm(val))

    normal = doc.styles["Normal"]
    normal.font.name = FONT_EN
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)

    lines = src.read_text(encoding="utf-8").splitlines()
    i = 0
    pending_bullet = None      # 连续的项目符号列表

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        i += 1

        stripped = line.strip()
        if stripped in ("---", ""):
            continue

        # ---- 一级：姓名 ----
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_runs(p, line[2:], size=17, base_bold=True, color=BLUE)
            _tight(p, 0, 2)
            continue

        # ---- 二级：章节（带下边框）----
        if line.startswith("## "):
            p = doc.add_paragraph()
            _add_runs(p, line[3:], size=12, base_bold=True, color=BLUE)
            _bottom_border(p)
            _tight(p, 8, 4)
            continue

        # ---- 三级：项目标题 ----
        if line.startswith("### "):
            p = doc.add_paragraph()
            _add_runs(p, line[4:], size=11, base_bold=True)
            _tight(p, 6, 2)
            continue

        # ---- 引用块：小字灰色（口径说明 / 备注）----
        if line.startswith(">"):
            p = doc.add_paragraph()
            _add_runs(p, line.lstrip("> ").strip(), size=8.5, color=GREY)
            _tight(p, 0, 3)
            continue

        # ---- 列表：二级缩进 ----
        m = re.match(r"^(\s*)-\s+(.*)$", line)
        if m:
            indent, text = len(m.group(1)), m.group(2)
            style = "List Bullet 2" if indent >= 2 else "List Bullet"
            p = doc.add_paragraph(style=style)
            _add_runs(p, text, size=9.5 if indent >= 2 else 10)
            _tight(p, 0, 1)
            continue

        # ---- 普通段落（**加粗开头**的多为字段行）----
        p = doc.add_paragraph()
        _add_runs(p, stripped, size=10)
        _tight(p, 0, 2)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(f"已生成: {out}")
    print(f"  段落数 {len(doc.paragraphs)} | 大小 {out.stat().st_size} 字节")


if __name__ == "__main__":
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SRC
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_OUT
    if not src.exists():
        sys.exit(f"[ERROR] 找不到源文件: {src}")
    build(src, out)
